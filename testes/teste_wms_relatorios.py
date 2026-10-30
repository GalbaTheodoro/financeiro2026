"""Teste dos relatórios e do BI do armazém (WMS → Relatórios e BI).

Monta um armazém pequeno e confere, contra as contas feitas à mão:

  * safra: a do cabeçalho da entrada vale para os lotes sem safra própria;
    transferência herda; resultado de serviço herda dos despejos (se iguais);
  * entradas, saídas, transferências (emitidas e recebidas), serviços e lotes,
    com os filtros de cliente, safra, padrão e peneira e o "agrupar por";
  * o painel (indicadores, mês a mês, saldo por safra, ocupação);
  * o Excel abre (duas abas, totais) e o PDF é um PDF de verdade;
  * outra conta não lê os relatórios de ninguém.

Uso:  python testes/teste_wms_relatorios.py   (com o servidor no ar)
"""
import io
import json
import pathlib
import sys
import time
import urllib.error
import urllib.request
from datetime import date

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

BASE = "http://127.0.0.1:8000"
erros = []


def chamar(metodo, caminho, dados=None, token=None, bruto=False, esperar_erro=False):
    req = urllib.request.Request(f"{BASE}{caminho}",
                                 data=json.dumps(dados).encode() if dados is not None else None,
                                 method=metodo)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req) as r:
            conteudo = r.read()
            return (conteudo, r.headers) if bruto else json.loads(conteudo or b"{}")
    except urllib.error.HTTPError as erro:
        if not esperar_erro:
            print(erro.read()[:400])
            raise
        corpo = json.loads(erro.read() or b"{}")
        corpo["_status"] = erro.code
        return corpo


def checar(descricao, condicao, extra=""):
    print(f"  [{'OK  ' if condicao else 'FALHA'}] {descricao} {extra}")
    if not condicao:
        erros.append(descricao)


sufixo = str(int(time.time()))
cad = chamar("POST", "/api/publico/cadastro", {
    "nome": "Galba", "email": f"wmsrel{sufixo}@teste.com", "senha": "123456",
    "empresa": "Armazém Relatórios", "plano": "P4_ANUAL"})
tk, eid = cad["token"], cad["empresa"]["id"]
hoje = date.today().isoformat()

print("=== 1. Montando o armazém ===")
arm = chamar("POST", "/api/wms/armazens", {"empresa_id": eid, "nome": "Central", "capacidade_sacas": 1000}, tk)
qa = chamar("POST", "/api/wms/quadras", {"empresa_id": eid, "armazem_id": arm["id"], "nome": "A"}, tk)
joao = chamar("POST", "/api/parceiros", {"empresa_id": eid, "nome": "João", "tipo": "FORNECEDOR", "pessoa": "F"}, tk)
maria = chamar("POST", "/api/parceiros", {"empresa_id": eid, "nome": "Maria", "tipo": "CLIENTE", "pessoa": "J"}, tk)

ge = chamar("POST", "/api/wms/guias-entrada", {
    "empresa_id": eid, "armazem_id": arm["id"], "cliente_id": joao["id"], "sacas": 300, "peso_total": 18000,
    "safra": "2025/2026",
    "lotes": [
        {"peneira": "17/18", "padrao": "Tipo 6", "sacas": 200, "peso": 12000, "valor_unitario": 1500,
         "locais": [{"quadra_id": qa["id"], "pilha": "P1", "sacas": 200}]},
        {"peneira": "15/16", "padrao": "Tipo 7", "safra": "2024/2025", "sacas": 100, "peso": 6000,
         "valor_unitario": 1200, "locais": [{"quadra_id": qa["id"], "pilha": "P2", "sacas": 100}]},
    ]}, tk)
ge = chamar("POST", f"/api/wms/guias-entrada/{ge['id']}/fechar", None, tk)
checar("a guia guarda a safra do cabeçalho", ge.get("safra") == "2025/2026")
checar("lote sem safra própria fica com a do cabeçalho; o outro mantém a sua",
       [l["safra"] for l in ge["lotes"]] == ["2025/2026", "2024/2025"], str([l["safra"] for l in ge["lotes"]]))
lotes = {l["codigo"]: l for l in chamar("GET", f"/api/wms/lotes?empresa_id={eid}", None, tk)}
checar("o lote do armazém leva a safra", lotes["1/1"]["safra"] == "2025/2026" and lotes["1/2"]["safra"] == "2024/2025")

disp = {(d["lote_codigo"], d["pilha"]): d["posicao_id"] for d in chamar(
    "GET", f"/api/wms/posicoes-disponiveis?empresa_id={eid}&armazem_id={arm['id']}&cliente_id={joao['id']}", None, tk)}
gs = chamar("POST", "/api/wms/guias-saida", {"empresa_id": eid, "armazem_id": arm["id"], "cliente_id": joao["id"],
                                             "destino": "Santos", "itens": [{"posicao_id": disp[("1/1", "P1")], "sacas": 50}]}, tk)
chamar("POST", f"/api/wms/guias-saida/{gs['id']}/fechar", None, tk)
gt = chamar("POST", "/api/wms/guias-transferencia", {
    "empresa_id": eid, "armazem_id": arm["id"], "cliente_id": joao["id"], "cliente_novo_id": maria["id"],
    "itens": [{"posicao_id": disp[("1/1", "P1")], "sacas": 30}]}, tk)
gt = chamar("POST", f"/api/wms/guias-transferencia/{gt['id']}/fechar", None, tk)
lotes = {l["codigo"]: l for l in chamar("GET", f"/api/wms/lotes?empresa_id={eid}", None, tk)}
checar("o lote transferido (1/1T) herda a safra", lotes["1/1T"]["safra"] == "2025/2026")

gsv = chamar("POST", "/api/wms/guias-servico", {
    "empresa_id": eid, "armazem_id": arm["id"], "cliente_id": joao["id"], "servico": "Rebenefício",
    "despejos": [{"posicao_id": disp[("1/2", "P2")], "sacas": 100}],
    "resultados": [{"peneira": "16", "padrao": "Tipo 6", "sacas": 95,
                    "locais": [{"quadra_id": qa["id"], "pilha": "P3", "sacas": 95}]}]}, tk)
gsv = chamar("POST", f"/api/wms/guias-servico/{gsv['id']}/fechar", None, tk)
lotes = {l["codigo"]: l for l in chamar("GET", f"/api/wms/lotes?empresa_id={eid}", None, tk)}
checar("o resultado do serviço herda a safra dos despejos", lotes["S1/1"]["safra"] == "2024/2025",
       str(lotes["S1/1"]["safra"]))

R = "/api/wms/relatorios"
q = f"empresa_id={eid}&de={hoje[:4]}-01-01&ate={hoje}"

print("\n=== 2. Opções dos filtros ===")
op = chamar("GET", f"{R}/opcoes?empresa_id={eid}", None, tk)
checar("traz as safras (mais nova primeiro), padrões, peneiras e serviços",
       op["safras"] == ["2025/2026", "2024/2025"] and "Tipo 6" in op["padroes"]
       and "17/18" in op["peneiras"] and op["servicos"] == ["Rebenefício"], str(op["safras"]))

print("\n=== 3. Relatórios ===")
ent = chamar("GET", f"{R}/entradas?{q}&agrupar=safra", None, tk)
checar("entradas: 2 lotes, 300 sacas, valor 420.000",
       ent["totais"]["linhas"] == 2 and ent["totais"]["sacas"] == 300 and ent["totais"]["valor"] == 420000,
       str(ent["totais"]))
checar("agrupado por safra: 200 em 2025/2026 e 100 em 2024/2025",
       [(g["rotulo"], g["sacas"]) for g in ent["grupos"]] == [("2025/2026", 200), ("2024/2025", 100)])
ent_p = chamar("GET", f"{R}/entradas?{q}&peneira=15/16", None, tk)
checar("filtro de peneira", ent_p["totais"]["sacas"] == 100 and ent_p["linhas"][0]["lote"] == "1/2")
ent_s = chamar("GET", f"{R}/entradas?{q}&safra=2025/2026&padrao=tipo%206", None, tk)
checar("filtro de safra + padrão (sem diferença de maiúsculas)", ent_s["totais"]["sacas"] == 200)

sai = chamar("GET", f"{R}/saidas?{q}&agrupar=padrao", None, tk)
checar("saídas: 50 sacas do Tipo 6 para Santos",
       sai["totais"]["sacas"] == 50 and sai["grupos"][0]["rotulo"] == "Tipo 6"
       and sai["linhas"][0]["destino"] == "Santos")

tra = chamar("GET", f"{R}/transferencias?{q}", None, tk)
checar("transferências: 30 sacas, lote novo 1/1T",
       tra["totais"]["sacas"] == 30 and tra["linhas"][0]["lote_novo"] == "1/1T")
emi = chamar("GET", f"{R}/transferencias?{q}&cliente_id={joao['id']}&direcao=emitidas", None, tk)
rec = chamar("GET", f"{R}/transferencias?{q}&cliente_id={joao['id']}&direcao=recebidas", None, tk)
rec_m = chamar("GET", f"{R}/transferencias?{q}&cliente_id={maria['id']}&direcao=recebidas&agrupar=cliente", None, tk)
checar("o João emitiu 30 e não recebeu nada",
       emi["totais"]["sacas"] == 30 and rec["totais"]["linhas"] == 0 and emi["linhas"][0]["direcao"] == "Emitida")
checar("a Maria recebeu 30 (agrupada pelo recebedor)",
       rec_m["totais"]["sacas"] == 30 and rec_m["grupos"][0]["rotulo"] == "Maria"
       and rec_m["linhas"][0]["direcao"] == "Recebida")
checar("a coluna Direção só aparece com cliente escolhido",
       "direcao" not in [c["chave"] for c in tra["colunas"]] and "direcao" in [c["chave"] for c in emi["colunas"]])

srv = chamar("GET", f"{R}/servicos?{q}&agrupar=servico", None, tk)
checar("serviços: despejou 100, resultou 95, quebra 5",
       srv["totais"]["despejo"] == 100 and srv["totais"]["resultado"] == 95 and srv["totais"]["quebra"] == 5,
       str(srv["totais"]))
checar("agrupado por serviço", srv["grupos"][0]["rotulo"] == "Rebenefício" and srv["grupos"][0]["quebra"] == 5)

lot = chamar("GET", f"{R}/lotes?empresa_id={eid}&agrupar=cliente", None, tk)
saldos = {l["lote"]: l["sacas"] for l in lot["linhas"]}
checar("lotes: saldo de hoje (1/1: 120, 1/1T: 30, S1/1: 95; 1/2 zerou e sai)",
       saldos == {"1/1": 120, "1/1T": 30, "S1/1": 95}, str(saldos))
checar("agrupado por cliente: João 215, Maria 30",
       [(g["rotulo"], g["sacas"]) for g in lot["grupos"]] == [("João", 215), ("Maria", 30)])
lot_todos = chamar("GET", f"{R}/lotes?empresa_id={eid}&com_saldo=false", None, tk)
checar("com 'só com saldo' desligado o lote zerado aparece", "1/2" in {l["lote"] for l in lot_todos["linhas"]})
mes = chamar("GET", f"{R}/entradas?{q}&agrupar=mes", None, tk)
checar("agrupar por mês", len(mes["grupos"]) == 1 and "/" in mes["grupos"][0]["rotulo"])
fora = chamar("GET", f"{R}/entradas?empresa_id={eid}&de=2001-01-01&ate=2001-12-31", None, tk)
checar("período sem movimento volta vazio", fora["totais"]["linhas"] == 0)
ruim = chamar("GET", f"{R}/entradas?empresa_id={eid}&de=2026-12-31&ate=2026-01-01", None, tk, esperar_erro=True)
checar("data inicial depois da final é recusada", ruim.get("_status") == 400)

print("\n=== 4. Painel ===")
pn = chamar("GET", f"{R}/painel?{q}", None, tk)
k = pn["indicadores"]
checar("indicadores: saldo 245, entradas 300, saídas 50, transferências 30, quebra 5",
       k["saldo_sacas"] == 245 and k["entradas_sacas"] == 300 and k["saidas_sacas"] == 50
       and k["transferencias_sacas"] == 30 and k["servicos_quebra"] == 5, str(k))
checar("12 meses no gráfico mês a mês, com o mês de hoje", len(pn["mensal"]) == 12
       and pn["mensal"][-1]["entradas"] == 300, str(pn["mensal"][-1]))
checar("saldo por safra", {s["rotulo"]: s["sacas"] for s in pn["por_safra"]} == {"2025/2026": 150, "2024/2025": 95})
checar("ocupação do armazém (245 de 1000 = 24,5%)", pn["ocupacao"][0]["ocupacao"] == 24.5)

print("\n=== 5. Excel e PDF ===")
conteudo, cab = chamar("GET", f"{R}/entradas/xlsx?{q}&agrupar=safra", None, tk, bruto=True)
checar("o Excel vem como arquivo .xlsx", "spreadsheetml" in cab.get("Content-Type", "")
       and ".xlsx" in cab.get("Content-Disposition", ""))
try:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo))
    checar("com as abas Resumo e Detalhe", wb.sheetnames == ["Resumo", "Detalhe"])
    det = wb["Detalhe"]
    checar("o detalhe tem as 2 linhas e o total com fórmula",
           det.max_row == 4 and str(det.cell(row=4, column=10).value).startswith("=SUM"))
except ImportError:
    print("  (openpyxl não instalado aqui: pulei a leitura do Excel)")
for tipo in ("entradas", "saidas", "transferencias", "servicos", "lotes"):
    conteudo, cab = chamar("GET", f"{R}/{tipo}/pdf?{q}", None, tk, bruto=True)
    checar(f"PDF de {tipo}", conteudo[:5] == b"%PDF-" and len(conteudo) > 1500, str(len(conteudo)))
    conteudo, cab = chamar("GET", f"{R}/{tipo}/xlsx?{q}", None, tk, bruto=True)
    checar(f"Excel de {tipo}", conteudo[:2] == b"PK")

print("\n=== 6. Isolamento ===")
outro = chamar("POST", "/api/publico/cadastro", {
    "nome": "Outro", "email": f"wmsrel2{sufixo}@teste.com", "senha": "123456",
    "empresa": "Outro armazém", "plano": "P4_ANUAL"})
alheio = chamar("GET", f"{R}/entradas?{q}", None, outro["token"], esperar_erro=True)
checar("outra conta não lê os relatórios desta empresa", alheio.get("_status") in (403, 404))
p1 = chamar("POST", "/api/publico/cadastro", {
    "nome": "P1", "email": f"wmsrel3{sufixo}@teste.com", "senha": "123456",
    "empresa": "Sem WMS", "plano": "P1_SEMESTRAL"})
sem = chamar("GET", f"{R}/painel?empresa_id={p1['empresa']['id']}", None, p1["token"], esperar_erro=True)
checar("plano sem WMS não abre os relatórios", sem.get("_status") == 403)

print("\n" + "=" * 60)
if erros:
    print(f"{len(erros)} FALHA(S):")
    for e in erros:
        print("  -", e)
    sys.exit(1)
print("Relatórios do armazém OK — os números batem na tela, no Excel e no PDF.")

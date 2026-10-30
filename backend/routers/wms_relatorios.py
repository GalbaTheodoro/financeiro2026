"""WMS — Relatórios e BI do armazém.

Uma tela com um **painel** (indicadores e gráficos) e cinco relatórios —
entradas, saídas, transferências (emitidas e recebidas), guias de serviço e
lotes (saldo de hoje) — todos com os mesmos filtros:

    período (de/até), armazém, cliente, safra, padrão e peneira

e um **agrupar por** (cliente, safra, padrão, peneira, armazém, mês e, nos
serviços, o tipo de serviço). Cada relatório sai na tela, em **Excel** (.xlsx,
com uma aba de resumo e outra de detalhe) e em **PDF** (A4 deitado, com o
gráfico do agrupamento, o resumo e o detalhe).

Só entram guias **fechadas** — rascunho e cancelada não movimentam o armazém.
O relatório de lotes é a fotografia de hoje: uma linha por lote e armazém com
saldo (o lote removido em parte para outro armazém aparece nos dois).

Rotas
-----
GET /api/wms/relatorios/opcoes             safras, padrões e peneiras existentes
GET /api/wms/relatorios/painel             indicadores e séries do painel
GET /api/wms/relatorios/{tipo}             linhas, grupos e totais
GET /api/wms/relatorios/{tipo}/xlsx        o mesmo, em Excel
GET /api/wms/relatorios/{tipo}/pdf         o mesmo, em PDF

``tipo``: entradas | saidas | transferencias | servicos | lotes
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import acesso_liberado, exigir_modulo, validar_empresa
from ..models import (
    Armazem,
    Empresa,
    GuiaEntrada,
    GuiaEntradaLote,
    GuiaSaida,
    GuiaSaidaItem,
    GuiaServico,
    GuiaServicoDespejo,
    GuiaServicoResultado,
    GuiaTransferencia,
    GuiaTransferenciaItem,
    LoteArmazem,
    Parceiro,
    Usuario,
)
from ..utils import parse_data
from .wms import TOLERANCIA, _somar_saldos

router = APIRouter(prefix="/api/wms/relatorios", tags=["wms-relatorios"],
                   dependencies=[Depends(exigir_modulo("WMS"))])

TIPOS = {
    "entradas": "Entradas",
    "saidas": "Saídas",
    "transferencias": "Transferências",
    "servicos": "Guias de serviço",
    "lotes": "Lotes — saldo de hoje",
}

AGRUPAMENTOS = {
    "cliente": "Cliente",
    "safra": "Safra",
    "padrao": "Padrão",
    "peneira": "Peneira",
    "armazem": "Armazém",
    "mes": "Mês",
    "servico": "Serviço",
}

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

# colunas de cada relatório: (chave, título, tipo) — tipo: t texto, d data, n número, m dinheiro
COLUNAS = {
    "entradas": [
        ("data", "Data", "d"), ("guia", "Guia", "t"), ("cliente", "Cliente", "t"),
        ("cliente_origem", "Origem", "t"), ("armazem", "Armazém", "t"), ("lote", "Lote", "t"),
        ("safra", "Safra", "t"), ("peneira", "Peneira", "t"), ("padrao", "Padrão", "t"),
        ("sacas", "Sacas", "n"), ("peso", "Peso (kg)", "n"), ("valor", "Valor (R$)", "m"),
    ],
    "saidas": [
        ("data", "Data", "d"), ("guia", "Guia", "t"), ("cliente", "Cliente", "t"),
        ("destino", "Destino", "t"), ("contrato", "Contrato", "t"), ("armazem", "Armazém", "t"),
        ("lote", "Lote", "t"), ("safra", "Safra", "t"), ("peneira", "Peneira", "t"),
        ("padrao", "Padrão", "t"), ("sacas", "Sacas", "n"), ("peso", "Peso (kg)", "n"),
    ],
    "transferencias": [
        ("data", "Data", "d"), ("guia", "Guia", "t"), ("direcao", "Direção", "t"),
        ("cliente", "Cedente", "t"), ("cliente_destino", "Recebedor", "t"),
        ("armazem", "Armazém", "t"), ("lote", "Lote origem", "t"), ("lote_novo", "Lote novo", "t"),
        ("safra", "Safra", "t"), ("peneira", "Peneira", "t"), ("padrao", "Padrão", "t"),
        ("sacas", "Sacas", "n"), ("peso", "Peso (kg)", "n"),
    ],
    "servicos": [
        ("data", "Data", "d"), ("guia", "Guia", "t"), ("servico", "Serviço", "t"),
        ("cliente", "Cliente", "t"), ("armazem", "Armazém", "t"), ("movimento", "Movimento", "t"),
        ("lote", "Lote", "t"), ("safra", "Safra", "t"), ("peneira", "Peneira", "t"),
        ("padrao", "Padrão", "t"), ("sacas", "Sacas", "n"), ("peso", "Peso (kg)", "n"),
    ],
    "lotes": [
        ("data", "Entrada", "d"), ("lote", "Lote", "t"), ("cliente", "Cliente", "t"),
        ("armazem", "Armazém", "t"), ("origem", "Origem", "t"), ("safra", "Safra", "t"),
        ("peneira", "Peneira", "t"), ("padrao", "Padrão", "t"), ("sacas_entrada", "Sacas entrada", "n"),
        ("sacas", "Saldo sacas", "n"), ("peso", "Saldo peso (kg)", "n"),
    ],
}


# --------------------------------------------------------------------------- #
# Filtros
# --------------------------------------------------------------------------- #
class Filtros:
    def __init__(self, empresa_id: int, de: str | None, ate: str | None,
                 armazem_id: int | None, cliente_id: int | None, safra: str | None,
                 padrao: str | None, peneira: str | None, agrupar: str | None,
                 direcao: str | None, servico: str | None, com_saldo: bool = True):
        self.empresa_id = empresa_id
        hoje = date.today()
        self.ate = parse_data(ate) or hoje
        self.de = parse_data(de) or date(self.ate.year, 1, 1)
        if self.de > self.ate:
            raise HTTPException(400, "A data inicial é depois da final.")
        self.armazem_id = armazem_id or None
        self.cliente_id = cliente_id or None
        self.safra = _limpo(safra)
        self.padrao = _limpo(padrao)
        self.peneira = _limpo(peneira)
        self.agrupar = agrupar if agrupar in AGRUPAMENTOS else "cliente"
        self.direcao = direcao if direcao in ("emitidas", "recebidas") else "todas"
        self.servico = _limpo(servico)
        self.com_saldo = com_saldo

    def atributos_ok(self, safra, padrao, peneira) -> bool:
        return ((not self.safra or _limpo(safra) == self.safra)
                and (not self.padrao or _limpo(padrao) == self.padrao)
                and (not self.peneira or _limpo(peneira) == self.peneira))

    def no_periodo(self, d: date | None) -> bool:
        return d is not None and self.de <= d <= self.ate


def _limpo(texto) -> str:
    return " ".join(str(texto or "").split()).upper()


def _filtros(empresa_id: int, de: str | None = None, ate: str | None = None,
             armazem_id: int | None = None, cliente_id: int | None = None,
             safra: str | None = None, padrao: str | None = None, peneira: str | None = None,
             agrupar: str | None = None, direcao: str | None = None,
             servico: str | None = None, com_saldo: bool = True) -> Filtros:
    return Filtros(empresa_id, de, ate, armazem_id, cliente_id, safra, padrao, peneira,
                   agrupar, direcao, servico, com_saldo)


def _nomes(db: Session, modelo, empresa_id: int) -> dict[int, str]:
    campo = "nome"
    return {o.id: getattr(o, campo) for o in db.query(modelo).filter(modelo.empresa_id == empresa_id)}


def _f(v) -> float:
    return round(float(v or 0), 2)


# --------------------------------------------------------------------------- #
# As linhas de cada relatório
# --------------------------------------------------------------------------- #
def linhas_entradas(db: Session, f: Filtros, nomes) -> list[dict]:
    q = (db.query(GuiaEntradaLote, GuiaEntrada).join(GuiaEntrada, GuiaEntradaLote.guia_id == GuiaEntrada.id)
         .filter(GuiaEntrada.empresa_id == f.empresa_id, GuiaEntrada.status == "FECHADA",
                 GuiaEntrada.data >= f.de, GuiaEntrada.data <= f.ate))
    if f.armazem_id:
        q = q.filter(GuiaEntrada.armazem_id == f.armazem_id)
    if f.cliente_id:
        q = q.filter(GuiaEntrada.cliente_id == f.cliente_id)
    codigos = {l.guia_lote_id: l.codigo for l in db.query(LoteArmazem).filter(
        LoteArmazem.empresa_id == f.empresa_id, LoteArmazem.guia_lote_id.isnot(None))}
    saida = []
    for l, g in q.all():
        safra = l.safra or g.safra
        if not f.atributos_ok(safra, l.padrao, l.peneira):
            continue
        saida.append({
            "data": g.data, "guia": str(g.numero), "cliente_id": g.cliente_id,
            "cliente": nomes["parceiro"].get(g.cliente_id, "-"),
            "cliente_origem": nomes["parceiro"].get(g.cliente_origem_id, "") if g.cliente_origem_id else "",
            "armazem_id": g.armazem_id, "armazem": nomes["armazem"].get(g.armazem_id, "-"),
            "lote": codigos.get(l.id) or f"{g.numero}/{l.sequencia}",
            "safra": safra or "", "peneira": l.peneira or "", "padrao": l.padrao or "",
            "sacas": _f(l.sacas), "peso": _f(l.peso), "valor": _f(l.valor_total),
        })
    return saida


def linhas_saidas(db: Session, f: Filtros, nomes) -> list[dict]:
    q = (db.query(GuiaSaidaItem, GuiaSaida).join(GuiaSaida, GuiaSaidaItem.guia_id == GuiaSaida.id)
         .filter(GuiaSaida.empresa_id == f.empresa_id, GuiaSaida.status == "FECHADA",
                 GuiaSaida.data >= f.de, GuiaSaida.data <= f.ate))
    if f.armazem_id:
        q = q.filter(GuiaSaida.armazem_id == f.armazem_id)
    if f.cliente_id:
        q = q.filter(GuiaSaida.cliente_id == f.cliente_id)
    saida = []
    for i, g in q.all():
        lote = i.lote
        if not lote or not f.atributos_ok(lote.safra, lote.padrao, lote.peneira):
            continue
        destino = " · ".join(x for x in (
            nomes["parceiro"].get(g.destinatario_id, "") if g.destinatario_id else "", g.destino or "") if x)
        saida.append({
            "data": g.data, "guia": str(g.numero), "cliente_id": g.cliente_id,
            "cliente": nomes["parceiro"].get(g.cliente_id, "-"), "destino": destino,
            "contrato": str(g.contrato.numero) if g.contrato else "",
            "armazem_id": g.armazem_id, "armazem": nomes["armazem"].get(g.armazem_id, "-"),
            "lote": lote.codigo, "safra": lote.safra or "", "peneira": lote.peneira or "",
            "padrao": lote.padrao or "", "sacas": _f(i.sacas), "peso": _f(i.peso),
        })
    return saida


def linhas_transferencias(db: Session, f: Filtros, nomes) -> list[dict]:
    q = (db.query(GuiaTransferenciaItem, GuiaTransferencia)
         .join(GuiaTransferencia, GuiaTransferenciaItem.guia_id == GuiaTransferencia.id)
         .filter(GuiaTransferencia.empresa_id == f.empresa_id, GuiaTransferencia.status == "FECHADA",
                 GuiaTransferencia.data >= f.de, GuiaTransferencia.data <= f.ate))
    if f.armazem_id:
        q = q.filter(GuiaTransferencia.armazem_id == f.armazem_id)
    codigos = {l.id: l.codigo for l in db.query(LoteArmazem).filter(LoteArmazem.empresa_id == f.empresa_id)}
    saida = []
    for i, g in q.all():
        lote = i.lote
        if not lote or not f.atributos_ok(lote.safra, lote.padrao, lote.peneira):
            continue
        direcao = ""
        if f.cliente_id:
            emitida, recebida = g.cliente_id == f.cliente_id, g.cliente_novo_id == f.cliente_id
            if not (emitida or recebida):
                continue
            if f.direcao == "emitidas" and not emitida:
                continue
            if f.direcao == "recebidas" and not recebida:
                continue
            direcao = "Emitida" if emitida else "Recebida"
        # no "por cliente", a transferência recebida conta para quem recebeu
        dono = g.cliente_novo_id if f.direcao == "recebidas" else g.cliente_id
        saida.append({
            "data": g.data, "guia": str(g.numero), "direcao": direcao,
            "cliente_id": dono, "cliente_agrupado": nomes["parceiro"].get(dono, "-"),
            "cliente": nomes["parceiro"].get(g.cliente_id, "-"),
            "cliente_destino": nomes["parceiro"].get(g.cliente_novo_id, "-"),
            "armazem_id": g.armazem_id, "armazem": nomes["armazem"].get(g.armazem_id, "-"),
            "lote": lote.codigo, "lote_novo": codigos.get(i.lote_novo_id, ""),
            "safra": lote.safra or "", "peneira": lote.peneira or "", "padrao": lote.padrao or "",
            "sacas": _f(i.sacas), "peso": _f(i.peso),
        })
    return saida


def linhas_servicos(db: Session, f: Filtros, nomes) -> list[dict]:
    q = db.query(GuiaServico).filter(
        GuiaServico.empresa_id == f.empresa_id, GuiaServico.status == "FECHADA",
        GuiaServico.data >= f.de, GuiaServico.data <= f.ate)
    if f.armazem_id:
        q = q.filter(GuiaServico.armazem_id == f.armazem_id)
    if f.cliente_id:
        q = q.filter(GuiaServico.cliente_id == f.cliente_id)
    saida = []
    for g in q.order_by(GuiaServico.data, GuiaServico.numero).all():
        if f.servico and _limpo(g.servico) != f.servico:
            continue
        base = {"data": g.data, "guia": str(g.numero), "servico": g.servico or "",
                "cliente_id": g.cliente_id, "cliente": nomes["parceiro"].get(g.cliente_id, "-"),
                "armazem_id": g.armazem_id, "armazem": nomes["armazem"].get(g.armazem_id, "-")}
        for d in g.despejos:
            lote = d.lote
            if not lote or not f.atributos_ok(lote.safra, lote.padrao, lote.peneira):
                continue
            saida.append(dict(base, movimento="Despejo", lote=lote.codigo, safra=lote.safra or "",
                              peneira=lote.peneira or "", padrao=lote.padrao or "",
                              sacas=_f(d.sacas), peso=_f(d.peso)))
        for r in g.resultados:
            gerado = r.lote_gerado
            safra = r.safra or (gerado.safra if gerado else None)
            if not f.atributos_ok(safra, r.padrao, r.peneira):
                continue
            saida.append(dict(base, movimento="Resultado", lote=gerado.codigo if gerado else "",
                              safra=safra or "", peneira=r.peneira or "", padrao=r.padrao or "",
                              sacas=_f(r.sacas), peso=_f(r.peso)))
    return saida


def linhas_lotes(db: Session, f: Filtros, nomes) -> list[dict]:
    """Uma linha por lote e armazém, com o saldo de hoje."""
    saldos = _somar_saldos(db, f.empresa_id)
    lotes: dict[tuple, LoteArmazem] = {}
    for l in db.query(LoteArmazem).filter(LoteArmazem.empresa_id == f.empresa_id):
        lotes.setdefault((l.depositante_id, l.produto_id, l.codigo.upper()), l)
    codigos = {l.id: l.codigo for l in lotes.values()}
    saida = []
    for (armazem_id, depositante_id, produto_id, codigo), v in saldos.items():
        lote = lotes.get((depositante_id, produto_id, codigo))
        if lote is None:
            continue
        if f.com_saldo and v["sacas"] <= TOLERANCIA:
            continue
        if f.armazem_id and armazem_id != f.armazem_id:
            continue
        if f.cliente_id and depositante_id != f.cliente_id:
            continue
        if not f.atributos_ok(lote.safra, lote.padrao, lote.peneira):
            continue
        origem = (f"transferido do {codigos.get(lote.lote_origem_id, '')}" if lote.lote_origem_id
                  else "resultado de serviço" if lote.guia_servico_id else "guia de entrada")
        saida.append({
            "data": lote.data_entrada, "lote": lote.codigo, "cliente_id": depositante_id,
            "cliente": nomes["parceiro"].get(depositante_id, "-"),
            "armazem_id": armazem_id, "armazem": nomes["armazem"].get(armazem_id, "-"),
            "origem": origem, "safra": lote.safra or "", "peneira": lote.peneira or "",
            "padrao": lote.padrao or "", "sacas_entrada": _f(lote.sacas),
            "sacas": round(v["sacas"], 2), "peso": round(v["peso_kg"], 2),
        })
    saida.sort(key=lambda x: (x["cliente"], x["data"] or date.min, x["lote"]))
    return saida


GERADORES = {
    "entradas": linhas_entradas,
    "saidas": linhas_saidas,
    "transferencias": linhas_transferencias,
    "servicos": linhas_servicos,
    "lotes": linhas_lotes,
}


# --------------------------------------------------------------------------- #
# Agrupar
# --------------------------------------------------------------------------- #
def _chave_grupo(linha: dict, agrupar: str) -> tuple[str, str]:
    """(chave para ordenar/juntar, rótulo)."""
    if agrupar == "mes":
        d = linha.get("data")
        if not d:
            return ("0000-00", "sem data")
        return (f"{d.year}-{d.month:02d}", f"{MESES[d.month - 1]}/{d.year}")
    if agrupar == "cliente":
        rotulo = linha.get("cliente_agrupado") or linha.get("cliente") or "-"
        return (rotulo.upper(), rotulo)
    valor = linha.get(agrupar) or ""
    rotulo = valor or "(não informado)"
    return (_limpo(valor) or "~", rotulo)


def agrupar(linhas: list[dict], agrupar_por: str, tipo: str) -> list[dict]:
    grupos: dict[str, dict] = {}
    for l in linhas:
        chave, rotulo = _chave_grupo(l, agrupar_por)
        g = grupos.setdefault(chave, {"chave": chave, "rotulo": rotulo, "linhas": 0,
                                      "sacas": 0.0, "peso": 0.0, "valor": 0.0,
                                      "despejo": 0.0, "resultado": 0.0})
        g["linhas"] += 1
        g["sacas"] += l.get("sacas", 0)
        g["peso"] += l.get("peso", 0)
        g["valor"] += l.get("valor", 0) or 0
        if tipo == "servicos":
            g["despejo" if l.get("movimento") == "Despejo" else "resultado"] += l.get("sacas", 0)
    lista = list(grupos.values())
    for g in lista:
        for k in ("sacas", "peso", "valor", "despejo", "resultado"):
            g[k] = round(g[k], 2)
        if tipo == "servicos":
            g["quebra"] = round(g["despejo"] - g["resultado"], 2)
            g["sacas"] = g["despejo"]          # o tamanho do grupo é o que entrou no serviço
    if agrupar_por == "mes":
        lista.sort(key=lambda g: g["chave"])
    else:
        lista.sort(key=lambda g: -g["sacas"])
    return lista


def totais(linhas: list[dict], tipo: str) -> dict:
    t = {"linhas": len(linhas),
         "sacas": round(sum(l.get("sacas", 0) for l in linhas), 2),
         "peso": round(sum(l.get("peso", 0) for l in linhas), 2),
         "valor": round(sum(l.get("valor", 0) or 0 for l in linhas), 2),
         "guias": len({l.get("guia") for l in linhas if l.get("guia")}),
         "clientes": len({l.get("cliente_id") for l in linhas if l.get("cliente_id")}),
         "lotes": len({l.get("lote") for l in linhas if l.get("lote")})}
    if tipo == "servicos":
        t["despejo"] = round(sum(l["sacas"] for l in linhas if l["movimento"] == "Despejo"), 2)
        t["resultado"] = round(sum(l["sacas"] for l in linhas if l["movimento"] == "Resultado"), 2)
        t["quebra"] = round(t["despejo"] - t["resultado"], 2)
        t["sacas"] = t["despejo"]
    return t


def montar(db: Session, tipo: str, f: Filtros) -> dict:
    if tipo not in GERADORES:
        raise HTTPException(404, "Relatório não encontrado.")
    nomes = {"parceiro": _nomes(db, Parceiro, f.empresa_id), "armazem": _nomes(db, Armazem, f.empresa_id)}
    linhas = GERADORES[tipo](db, f, nomes)
    if tipo != "lotes":
        linhas.sort(key=lambda x: (x["data"] or date.min, int(x["guia"]) if str(x.get("guia", "")).isdigit() else 0))
    agrupar_por = f.agrupar if not (f.agrupar == "servico" and tipo != "servicos") else "cliente"
    return {
        "tipo": tipo, "titulo": TIPOS[tipo],
        "colunas": [{"chave": c, "titulo": t, "formato": k} for c, t, k in COLUNAS[tipo]
                    if not (c == "direcao" and not f.cliente_id)],
        "agrupar": agrupar_por, "agrupar_rotulo": AGRUPAMENTOS[agrupar_por],
        "grupos": agrupar(linhas, agrupar_por, tipo),
        "totais": totais(linhas, tipo),
        "linhas": linhas,
        "filtros": _descrever(db, tipo, f, nomes),
    }


def _descrever(db: Session, tipo: str, f: Filtros, nomes) -> str:
    partes = []
    if tipo != "lotes":
        partes.append(f"Período {f.de:%d/%m/%Y} a {f.ate:%d/%m/%Y}")
    else:
        partes.append(f"Posição em {date.today():%d/%m/%Y}")
    if f.armazem_id:
        partes.append(f"Armazém {nomes['armazem'].get(f.armazem_id, '-')}")
    if f.cliente_id:
        partes.append(f"Cliente {nomes['parceiro'].get(f.cliente_id, '-')}")
    if tipo == "transferencias" and f.cliente_id and f.direcao != "todas":
        partes.append("só emitidas" if f.direcao == "emitidas" else "só recebidas")
    for rotulo, valor in (("Safra", f.safra), ("Padrão", f.padrao), ("Peneira", f.peneira),
                          ("Serviço", f.servico if tipo == "servicos" else "")):
        if valor:
            partes.append(f"{rotulo} {valor}")
    return " · ".join(partes)


# --------------------------------------------------------------------------- #
# Rotas
# --------------------------------------------------------------------------- #
@router.get("/opcoes")
def opcoes(empresa_id: int, db: Session = Depends(get_db),
           usuario: Usuario = Depends(acesso_liberado)):
    """As safras, padrões, peneiras e serviços que já existem — para os filtros."""
    validar_empresa(db, empresa_id, usuario)
    safras, padroes, peneiras = set(), set(), set()
    for l in db.query(LoteArmazem).filter(LoteArmazem.empresa_id == empresa_id):
        if l.safra:
            safras.add(l.safra.strip())
        if l.padrao:
            padroes.add(" ".join(l.padrao.split()))
        if l.peneira:
            peneiras.add(l.peneira.strip())
    servicos = {(g.servico or "").strip() for g in db.query(GuiaServico).filter(
        GuiaServico.empresa_id == empresa_id) if (g.servico or "").strip()}
    juntar = lambda c: sorted({v.upper(): v for v in c}.values(), key=str.upper)  # noqa: E731
    return {"safras": sorted(juntar(safras), reverse=True), "padroes": juntar(padroes),
            "peneiras": juntar(peneiras), "servicos": juntar(servicos),
            "agrupamentos": AGRUPAMENTOS, "tipos": TIPOS}


@router.get("/painel")
def painel(empresa_id: int, de: str | None = None, ate: str | None = None,
           armazem_id: int | None = None, cliente_id: int | None = None,
           safra: str | None = None, padrao: str | None = None, peneira: str | None = None,
           db: Session = Depends(get_db), usuario: Usuario = Depends(acesso_liberado)):
    """Indicadores do período e a fotografia do saldo de hoje."""
    validar_empresa(db, empresa_id, usuario)
    f = _filtros(empresa_id, de, ate, armazem_id, cliente_id, safra, padrao, peneira)
    nomes = {"parceiro": _nomes(db, Parceiro, empresa_id), "armazem": _nomes(db, Armazem, empresa_id)}
    ent = linhas_entradas(db, f, nomes)
    sai = linhas_saidas(db, f, nomes)
    tra = linhas_transferencias(db, f, nomes)
    srv = linhas_servicos(db, f, nomes)
    lot = linhas_lotes(db, f, nomes)

    # movimento mês a mês: os 12 meses que terminam no "até"
    f12 = _filtros(empresa_id, (date(f.ate.year, f.ate.month, 1) - timedelta(days=334)).replace(day=1)
                   .isoformat(), f.ate.isoformat(), armazem_id, cliente_id, safra, padrao, peneira)
    meses = []
    cursor = f12.de
    while cursor <= f12.ate:
        meses.append((cursor.year, cursor.month))
        cursor = (cursor.replace(day=28) + timedelta(days=4)).replace(day=1)
    por_mes = {m: {"entradas": 0.0, "saidas": 0.0, "transferencias": 0.0} for m in meses}
    for chave, linhas in (("entradas", linhas_entradas(db, f12, nomes)),
                          ("saidas", linhas_saidas(db, f12, nomes)),
                          ("transferencias", linhas_transferencias(db, f12, nomes))):
        for l in linhas:
            m = (l["data"].year, l["data"].month)
            if m in por_mes:
                por_mes[m][chave] += l["sacas"]
    mensal = [{"rotulo": f"{MESES[m - 1]}/{str(a)[2:]}", **{k: round(v, 2) for k, v in por_mes[(a, m)].items()}}
              for a, m in meses]

    def top(agrupar_por, n=10):
        grupos = agrupar(lot, agrupar_por, "lotes")
        return [{"rotulo": g["rotulo"], "sacas": g["sacas"], "peso": g["peso"], "lotes": g["linhas"]}
                for g in grupos[:n]]

    ocupacao = []
    saldo_armazem = defaultdict(float)
    for l in lot:
        saldo_armazem[l["armazem_id"]] += l["sacas"]
    for a in db.query(Armazem).filter(Armazem.empresa_id == empresa_id, Armazem.ativo.is_(True)):
        if armazem_id and a.id != armazem_id:
            continue
        cap = float(a.capacidade_sacas or 0)
        ocupacao.append({"rotulo": a.nome, "sacas": round(saldo_armazem.get(a.id, 0.0), 2),
                         "capacidade": round(cap, 2),
                         "ocupacao": round(saldo_armazem.get(a.id, 0.0) / cap * 100, 1) if cap else None})

    t_srv = totais(srv, "servicos")
    return {
        "filtros": _descrever(db, "painel", f, nomes),
        "indicadores": {
            "saldo_sacas": round(sum(l["sacas"] for l in lot), 2),
            "saldo_peso": round(sum(l["peso"] for l in lot), 2),
            "lotes_com_saldo": len(lot),
            "clientes_com_saldo": len({l["cliente_id"] for l in lot}),
            "entradas_sacas": round(sum(l["sacas"] for l in ent), 2),
            "entradas_guias": len({l["guia"] for l in ent}),
            "saidas_sacas": round(sum(l["sacas"] for l in sai), 2),
            "saidas_guias": len({l["guia"] for l in sai}),
            "transferencias_sacas": round(sum(l["sacas"] for l in tra), 2),
            "transferencias_guias": len({l["guia"] for l in tra}),
            "servicos_guias": t_srv["guias"],
            "servicos_despejo": t_srv["despejo"],
            "servicos_resultado": t_srv["resultado"],
            "servicos_quebra": t_srv["quebra"],
        },
        "mensal": mensal,
        "por_cliente": top("cliente"),
        "por_safra": top("safra", 8),
        "por_padrao": top("padrao"),
        "por_peneira": top("peneira", 12),
        "ocupacao": ocupacao,
    }


def _rota_filtros(tipo: str, empresa_id: int, de, ate, armazem_id, cliente_id, safra, padrao,
                  peneira, agrupar_por, direcao, servico, com_saldo, db, usuario) -> dict:
    validar_empresa(db, empresa_id, usuario)
    f = _filtros(empresa_id, de, ate, armazem_id, cliente_id, safra, padrao, peneira,
                 agrupar_por, direcao, servico, com_saldo)
    return montar(db, tipo, f)


@router.get("/{tipo}")
def relatorio(tipo: str, empresa_id: int, de: str | None = None, ate: str | None = None,
              armazem_id: int | None = None, cliente_id: int | None = None,
              safra: str | None = None, padrao: str | None = None, peneira: str | None = None,
              agrupar: str | None = None, direcao: str | None = None, servico: str | None = None,
              com_saldo: bool = True, db: Session = Depends(get_db),
              usuario: Usuario = Depends(acesso_liberado)):
    dados = _rota_filtros(tipo, empresa_id, de, ate, armazem_id, cliente_id, safra, padrao,
                          peneira, agrupar, direcao, servico, com_saldo, db, usuario)
    for l in dados["linhas"]:
        if isinstance(l.get("data"), date):
            l["data"] = l["data"].isoformat()
    return dados


def _nome_arquivo(tipo: str, ext: str) -> str:
    return f"armazem-{tipo}-{date.today():%Y-%m-%d}.{ext}"


@router.get("/{tipo}/xlsx")
def relatorio_xlsx(tipo: str, empresa_id: int, de: str | None = None, ate: str | None = None,
                   armazem_id: int | None = None, cliente_id: int | None = None,
                   safra: str | None = None, padrao: str | None = None, peneira: str | None = None,
                   agrupar: str | None = None, direcao: str | None = None, servico: str | None = None,
                   com_saldo: bool = True, db: Session = Depends(get_db),
                   usuario: Usuario = Depends(acesso_liberado)):
    dados = _rota_filtros(tipo, empresa_id, de, ate, armazem_id, cliente_id, safra, padrao,
                          peneira, agrupar, direcao, servico, com_saldo, db, usuario)
    empresa = db.get(Empresa, empresa_id)
    conteudo = gerar_xlsx(dados, _nome_empresa(empresa))
    return StreamingResponse(
        io.BytesIO(conteudo),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_nome_arquivo(tipo, "xlsx")}"'})


@router.get("/{tipo}/pdf")
def relatorio_pdf(tipo: str, empresa_id: int, de: str | None = None, ate: str | None = None,
                  armazem_id: int | None = None, cliente_id: int | None = None,
                  safra: str | None = None, padrao: str | None = None, peneira: str | None = None,
                  agrupar: str | None = None, direcao: str | None = None, servico: str | None = None,
                  com_saldo: bool = True, db: Session = Depends(get_db),
                  usuario: Usuario = Depends(acesso_liberado)):
    dados = _rota_filtros(tipo, empresa_id, de, ate, armazem_id, cliente_id, safra, padrao,
                          peneira, agrupar, direcao, servico, com_saldo, db, usuario)
    empresa = db.get(Empresa, empresa_id)
    conteudo = gerar_pdf(dados, _nome_empresa(empresa))
    return StreamingResponse(
        io.BytesIO(conteudo), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_nome_arquivo(tipo, "pdf")}"'})


def _nome_empresa(empresa) -> str:
    if not empresa:
        return ""
    return empresa.nome_fantasia or empresa.razao_social or ""


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _colunas_grupo(dados: dict) -> list[tuple[str, str, str]]:
    cols = [("rotulo", dados["agrupar_rotulo"], "t"), ("linhas", "Linhas", "i")]
    if dados["tipo"] == "servicos":
        cols += [("despejo", "Despejado (sc)", "n"), ("resultado", "Resultado (sc)", "n"),
                 ("quebra", "Quebra (sc)", "n")]
    else:
        cols += [("sacas", "Sacas", "n"), ("peso", "Peso (kg)", "n")]
        if dados["tipo"] == "entradas":
            cols.append(("valor", "Valor (R$)", "m"))
    cols.append(("pct", "% das sacas", "p"))
    return cols


def gerar_xlsx(dados: dict, empresa: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    verde = "15774A"
    cab_font = Font(bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor=verde)
    tot_fill = PatternFill("solid", fgColor="E8F4EE")
    fino = Side(style="thin", color="E2E8F0")
    borda = Border(bottom=fino)
    formatos = {"n": "#,##0.00", "m": "#,##0.00", "i": "0", "p": "0.0%", "d": "DD/MM/YYYY"}

    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    resumo["A1"] = f"{empresa} — Armazém: {dados['titulo']}" if empresa else f"Armazém: {dados['titulo']}"
    resumo["A1"].font = Font(bold=True, size=14, color="0E2B1C")
    resumo["A2"] = dados["filtros"]
    resumo["A2"].font = Font(italic=True, color="64748B")
    resumo["A3"] = f"Gerado em {datetime.now():%d/%m/%Y %H:%M} · agrupado por {dados['agrupar_rotulo'].lower()}"
    resumo["A3"].font = Font(color="64748B")

    cols = _colunas_grupo(dados)
    linha0 = 5
    for j, (_, titulo, _) in enumerate(cols, start=1):
        c = resumo.cell(row=linha0, column=j, value=titulo)
        c.font, c.fill = cab_font, cab_fill
        c.alignment = Alignment(horizontal="center" if j > 1 else "left")
    total_sacas = dados["totais"]["sacas"] or 0
    for i, g in enumerate(dados["grupos"], start=linha0 + 1):
        for j, (chave, _, fmt) in enumerate(cols, start=1):
            valor = (g["sacas"] / total_sacas if total_sacas else 0) if chave == "pct" else g.get(chave)
            c = resumo.cell(row=i, column=j, value=valor)
            c.border = borda
            if fmt in formatos:
                c.number_format = formatos[fmt]
    fim = linha0 + len(dados["grupos"])
    tl = fim + 1
    resumo.cell(row=tl, column=1, value="Total").font = Font(bold=True)
    for j, (chave, _, fmt) in enumerate(cols, start=1):
        cel = resumo.cell(row=tl, column=j)
        cel.fill = tot_fill
        if j == 1:
            continue
        letra = get_column_letter(j)
        if chave == "pct":
            cel.value = 1 if dados["grupos"] else 0
        elif fmt in ("n", "m", "i"):
            cel.value = f"=SUM({letra}{linha0 + 1}:{letra}{fim})" if dados["grupos"] else 0
        cel.number_format = formatos.get(fmt, "General")
        cel.font = Font(bold=True)
    resumo.column_dimensions["A"].width = 34
    for j in range(2, len(cols) + 1):
        resumo.column_dimensions[get_column_letter(j)].width = 16
    resumo.freeze_panes = resumo.cell(row=linha0 + 1, column=1)

    # gráfico do agrupamento (até 15 grupos)
    if dados["grupos"]:
        n = min(len(dados["grupos"]), 15)
        col_valor = 3  # sacas ou despejado
        grafico = BarChart()
        grafico.type = "bar" if dados["agrupar"] != "mes" else "col"
        grafico.style = 10
        grafico.title = f"Sacas por {dados['agrupar_rotulo'].lower()}"
        grafico.legend = None
        grafico.add_data(Reference(resumo, min_col=col_valor, min_row=linha0, max_row=linha0 + n),
                         titles_from_data=True)
        grafico.set_categories(Reference(resumo, min_col=1, min_row=linha0 + 1, max_row=linha0 + n))
        grafico.height, grafico.width = 9, 18
        if grafico.type == "bar":
            grafico.y_axis.scaling.orientation = "minMax"
            grafico.x_axis.scaling.orientation = "maxMin"
        resumo.add_chart(grafico, f"{get_column_letter(len(cols) + 2)}{linha0}")

    detalhe = wb.create_sheet("Detalhe")
    colunas = dados["colunas"]
    for j, c in enumerate(colunas, start=1):
        cel = detalhe.cell(row=1, column=j, value=c["titulo"])
        cel.font, cel.fill = cab_font, cab_fill
    for i, l in enumerate(dados["linhas"], start=2):
        for j, c in enumerate(colunas, start=1):
            valor = l.get(c["chave"])
            if c["formato"] == "d" and isinstance(valor, str) and valor:
                valor = date.fromisoformat(valor)
            cel = detalhe.cell(row=i, column=j, value=valor)
            if c["formato"] in formatos:
                cel.number_format = formatos[c["formato"]]
    ultima = len(dados["linhas"]) + 1
    if dados["linhas"]:
        tl = ultima + 1
        detalhe.cell(row=tl, column=1, value="Total").font = Font(bold=True)
        for j, c in enumerate(colunas, start=1):
            detalhe.cell(row=tl, column=j).fill = tot_fill
            if c["formato"] in ("n", "m") and dados["tipo"] != "servicos":
                letra = get_column_letter(j)
                cel = detalhe.cell(row=tl, column=j, value=f"=SUM({letra}2:{letra}{ultima})")
                cel.number_format = formatos[c["formato"]]
                cel.font = Font(bold=True)
    larguras = {"t": 18, "d": 12, "n": 14, "m": 15}
    for j, c in enumerate(colunas, start=1):
        largura = larguras.get(c["formato"], 14)
        if c["chave"] in ("cliente", "cliente_destino", "cliente_origem", "destino", "padrao"):
            largura = 28
        detalhe.column_dimensions[get_column_letter(j)].width = largura
    detalhe.freeze_panes = "A2"
    detalhe.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(ultima, 1)}"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _latin(texto) -> str:
    """As fontes padrão do PDF são latin-1: troca o que não cabe."""
    t = str(texto if texto is not None else "")
    for de, para in (("—", "-"), ("–", "-"), ("·", "|"), ("“", '"'), ("”", '"'), ("’", "'"), ("…", "...")):
        t = t.replace(de, para)
    return t.encode("latin-1", "replace").decode("latin-1")


def _num_br(v, casas=2) -> str:
    return f"{float(v or 0):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data_br(v) -> str:
    if isinstance(v, str) and v:
        v = date.fromisoformat(v)
    return v.strftime("%d/%m/%Y") if isinstance(v, date) else ""


def gerar_pdf(dados: dict, empresa: str) -> bytes:
    from fpdf import FPDF

    verde, mata, cinza = (21, 119, 74), (14, 43, 28), (100, 116, 139)
    titulo = f"Armazém - {dados['titulo']}"

    class Relatorio(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 13)
            self.set_text_color(*mata)
            self.cell(0, 7, _latin(titulo), new_x="LMARGIN", new_y="NEXT")
            self.set_font("Helvetica", "", 8.5)
            self.set_text_color(*cinza)
            self.cell(0, 4.5, _latin(" | ".join(x for x in (empresa, dados["filtros"]) if x)),
                      new_x="LMARGIN", new_y="NEXT")
            self.set_draw_color(*verde)
            self.set_line_width(0.6)
            self.line(self.l_margin, self.get_y() + 1.5, self.w - self.r_margin, self.get_y() + 1.5)
            self.ln(5)

        def footer(self):
            self.set_y(-10)
            self.set_font("Helvetica", "", 7.5)
            self.set_text_color(*cinza)
            self.cell(0, 5, _latin(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} - AgroDock"), align="L")
            self.cell(0, 5, f"Página {self.page_no()}/{{nb}}", align="R")

    pdf = Relatorio(orientation="L", unit="mm", format="A4")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(True, margin=14)
    pdf.set_margins(12, 10, 12)
    pdf.add_page()
    largura_util = pdf.w - pdf.l_margin - pdf.r_margin

    # indicadores
    t = dados["totais"]
    if dados["tipo"] == "servicos":
        kpis = [("Guias", str(t["guias"])), ("Despejado (sc)", _num_br(t["despejo"])),
                ("Resultado (sc)", _num_br(t["resultado"])), ("Quebra (sc)", _num_br(t["quebra"]))]
    else:
        kpis = [("Linhas", str(t["linhas"])), ("Sacas", _num_br(t["sacas"])),
                ("Peso (kg)", _num_br(t["peso"])), ("Clientes", str(t["clientes"]))]
        if dados["tipo"] == "entradas":
            kpis.append(("Valor (R$)", _num_br(t["valor"])))
        elif dados["tipo"] != "lotes":
            kpis.append(("Guias", str(t["guias"])))
        else:
            kpis.append(("Lotes", str(t["lotes"])))
    w = (largura_util - 3 * (len(kpis) - 1)) / len(kpis)
    y0 = pdf.get_y()
    for i, (rot, val) in enumerate(kpis):
        x = pdf.l_margin + i * (w + 3)
        pdf.set_fill_color(232, 244, 238)
        pdf.rect(x, y0, w, 14, style="F")
        pdf.set_fill_color(*verde)
        pdf.rect(x, y0, 1.2, 14, style="F")
        pdf.set_xy(x + 4, y0 + 2)
        pdf.set_font("Helvetica", "", 7.5)
        pdf.set_text_color(*cinza)
        pdf.cell(w - 6, 4, _latin(rot.upper()))
        pdf.set_xy(x + 4, y0 + 6.5)
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(*mata)
        pdf.cell(w - 6, 6, _latin(val))
    pdf.set_y(y0 + 19)

    # gráfico + resumo do agrupamento lado a lado
    grupos = dados["grupos"]
    if grupos:
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(*mata)
        pdf.cell(0, 6, _latin(f"Por {dados['agrupar_rotulo'].lower()}"), new_x="LMARGIN", new_y="NEXT")
        y_ini = pdf.get_y()
        mostrar = grupos[:12]
        maior = max((g["sacas"] for g in mostrar), default=0) or 1
        larg_graf = largura_util * 0.42
        rot_w = 42
        barra_max = larg_graf - rot_w - 22
        pdf.set_font("Helvetica", "", 7.5)
        for i, g in enumerate(mostrar):
            y = y_ini + i * 6
            pdf.set_xy(pdf.l_margin, y)
            pdf.set_text_color(64, 80, 70)
            pdf.cell(rot_w, 5, _latin(str(g["rotulo"])[:28]))
            comp = max(0.5, barra_max * (g["sacas"] / maior))
            pdf.set_fill_color(*verde)
            pdf.rect(pdf.l_margin + rot_w, y + 0.8, comp, 3.6, style="F")
            pdf.set_xy(pdf.l_margin + rot_w + comp + 1.5, y)
            pdf.set_text_color(*mata)
            pdf.cell(20, 5, _num_br(g["sacas"], 0))
        # tabela do resumo à direita
        cols = _colunas_grupo(dados)
        x_tab = pdf.l_margin + larg_graf + 6
        larg_tab = largura_util - larg_graf - 6
        pesos = [3] + [1.3] * (len(cols) - 1)
        ws = [larg_tab * p / sum(pesos) for p in pesos]
        pdf.set_xy(x_tab, y_ini)
        _cabecalho(pdf, [c[1] for c in cols], ws, x_tab, [c[2] for c in cols])
        total = t["sacas"] or 0
        for i, g in enumerate(grupos[:14]):
            if i % 2:
                pdf.set_fill_color(248, 250, 252)
            pdf.set_x(x_tab)
            pdf.set_font("Helvetica", "", 7.5)
            pdf.set_text_color(30, 41, 59)
            for (chave, _, fmt), larg in zip(cols, ws):
                if chave == "pct":
                    texto = _num_br((g["sacas"] / total * 100) if total else 0, 1) + "%"
                elif fmt == "i":
                    texto = str(g[chave])
                elif fmt in ("n", "m"):
                    texto = _num_br(g.get(chave, 0))
                else:
                    texto = str(g.get(chave, ""))[:32]
                pdf.cell(larg, 5, _latin(texto), align="L" if fmt == "t" else "R", fill=bool(i % 2))
            pdf.ln(5)
        if len(grupos) > 14:
            pdf.set_x(x_tab)
            pdf.set_font("Helvetica", "I", 7)
            pdf.set_text_color(*cinza)
            pdf.cell(larg_tab, 4.5, _latin(f"... e mais {len(grupos) - 14} (todos no Excel)"))
            pdf.ln(4.5)
        pdf.set_y(max(pdf.get_y(), y_ini + len(mostrar) * 6) + 6)

    # detalhe
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(*mata)
    pdf.cell(0, 6, "Detalhe", new_x="LMARGIN", new_y="NEXT")
    colunas = dados["colunas"]
    peso_col = {"t": 1.6, "d": 1.25, "n": 1.1, "m": 1.25}
    largos = {"cliente": 2.6, "cliente_destino": 2.4, "cliente_origem": 1.9, "destino": 2.2,
              "padrao": 2.1, "origem": 1.8, "servico": 1.5}
    pesos = [largos.get(c["chave"], peso_col.get(c["formato"], 1.2)) for c in colunas]
    ws = [largura_util * p / sum(pesos) for p in pesos]
    fmts = [c["formato"] for c in colunas]
    _cabecalho(pdf, [c["titulo"] for c in colunas], ws, pdf.l_margin, fmts)
    for i, l in enumerate(dados["linhas"]):
        if pdf.get_y() > pdf.h - 20:
            pdf.add_page()
            _cabecalho(pdf, [c["titulo"] for c in colunas], ws, pdf.l_margin, fmts)
        if i % 2:
            pdf.set_fill_color(248, 250, 252)
        pdf.set_font("Helvetica", "", 7.2)
        pdf.set_text_color(30, 41, 59)
        for c, larg in zip(colunas, ws):
            v = l.get(c["chave"])
            if c["formato"] == "d":
                texto = _data_br(v)
            elif c["formato"] in ("n", "m"):
                texto = _num_br(v)
            else:
                texto = str(v or "")
            maximo = max(4, int(larg / 1.45))
            if c["formato"] == "t" and len(texto) > maximo:
                texto = texto[:maximo - 1] + "."
            pdf.cell(larg, 4.8, _latin(texto), align="R" if c["formato"] in ("n", "m") else "L",
                     fill=bool(i % 2))
        pdf.ln(4.8)
    # total do detalhe
    if dados["linhas"]:
        pdf.set_fill_color(232, 244, 238)
        pdf.set_font("Helvetica", "B", 7.5)
        pdf.set_text_color(*mata)
        for j, (c, larg) in enumerate(zip(colunas, ws)):
            texto = ""
            if j == 0:
                texto = f"Total ({len(dados['linhas'])})"
            elif c["chave"] in ("sacas", "peso", "valor", "sacas_entrada") and dados["tipo"] != "servicos":
                texto = _num_br(sum(float(l.get(c["chave"]) or 0) for l in dados["linhas"]))
            pdf.cell(larg, 5.5, _latin(texto), align="R" if j else "L", fill=True)
        pdf.ln(5.5)
        if dados["tipo"] == "servicos":
            pdf.ln(2)
            pdf.set_font("Helvetica", "", 8)
            pdf.cell(0, 5, _latin(f"Despejado {_num_br(t['despejo'])} sc | Resultado "
                                  f"{_num_br(t['resultado'])} sc | Quebra {_num_br(t['quebra'])} sc"))
    else:
        pdf.set_font("Helvetica", "I", 9)
        pdf.set_text_color(*cinza)
        pdf.cell(0, 8, "Nenhum movimento com estes filtros.")
    return bytes(pdf.output())


def _cabecalho(pdf, titulos, larguras, x, formatos):
    pdf.set_x(x)
    pdf.set_fill_color(21, 119, 74)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 7.3)
    for titulo, larg, fmt in zip(titulos, larguras, formatos):
        pdf.cell(larg, 5.5, _latin(titulo), align="L" if fmt in ("t", "d") else "R", fill=True)
    pdf.ln(5.5)
